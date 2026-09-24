"""
extractor.py — Convierte los archivos de una carpeta de beca (PDF, DOCX, DOC,
HTML guardado, TXT) en un único documento Markdown con texto y tablas.

Las tablas se exportan en formato GFM (GitHub Flavored Markdown) para que
guion.py pueda linearizarlas después sin perder su significado.
"""

from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

EXTENSIONES = {".pdf", ".docx", ".doc", ".html", ".htm", ".txt"}
# Carpetas de assets de páginas web guardadas ("..._files") se ignoran
PATRON_ASSETS = re.compile(r"_files?$", re.IGNORECASE)


# ---------------------------------------------------------------- utilidades

def tabla_a_md(data: list[list], forzar_clave_valor: bool = False) -> str:
    """Convierte una tabla (lista de filas) a Markdown GFM.

    - Limpia celdas y descarta filas/columnas vacías.
    - Tablas de 2 columnas con pinta de ficha (campo/detalle) usan encabezado
      genérico para que TODAS las filas se conserven como datos.
    """
    filas = [
        [re.sub(r"\s+", " ", str(c or "")).strip() for c in fila]
        for fila in data
    ]
    filas = [f for f in filas if any(f)]
    if not filas:
        return ""

    ncols = max(len(f) for f in filas)
    filas = [f + [""] * (ncols - len(f)) for f in filas]

    # Quitar columnas completamente vacías
    cols_utiles = [i for i in range(ncols) if any(f[i] for f in filas)]
    filas = [[f[i] for i in cols_utiles] for f in filas]
    if not filas or not filas[0]:
        return ""
    ncols = len(filas[0])

    es_ficha = (
        forzar_clave_valor
        or (ncols == 2 and len(filas) >= 3
            and sum(1 for f in filas if 0 < len(f[0]) <= 60) >= len(filas) * 0.6)
    )

    if es_ficha:
        cabecera, cuerpo = ["Campo", "Detalle"], filas
    else:
        cabecera, cuerpo = filas[0], filas[1:]
        cabecera = [h if h else f"Columna {i + 1}" for i, h in enumerate(cabecera)]

    lineas = [
        "| " + " | ".join(cabecera) + " |",
        "|" + " --- |" * ncols,
    ]
    lineas += ["| " + " | ".join(f) + " |" for f in cuerpo]
    return "\n".join(lineas)


# ---------------------------------------------------------------------- PDF

def _pagina_a_md(page) -> str:
    """Modo básico por página: tablas detectadas + texto (sin duplicar celdas)."""
    partes, celdas = [], set()
    try:
        tabs = page.find_tables(snap_tolerance=3, join_tolerance=3)
    except Exception:
        tabs = None
    if tabs and tabs.tables:
        for t in tabs.tables:
            table_data = t.extract()
            md = tabla_a_md(table_data)
            if md:
                partes.append(md)
                for fila in table_data:
                    for c in fila:
                        for lin in str(c or "").split("\n"):
                            lin = lin.strip()
                            if lin:
                                celdas.add(lin)
    texto = page.get_text("text") or ""
    lineas = [
        l.rstrip() for l in texto.split("\n")
        if not l.strip() or l.strip() not in celdas
    ]
    partes.append("\n".join(lineas))
    return "\n\n".join(p for p in partes if p.strip())


def pdf_a_md(ruta: Path) -> str:
    import pymupdf

    doc = pymupdf.open(ruta)
    # 1) Intento con pymupdf4llm (mejor detección de tablas y orden de lectura)
    try:
        import pymupdf4llm

        md = pymupdf4llm.to_markdown(doc, show_progress=False)
        if md and len(md.strip()) > 40:
            return md
    except Exception:
        # p. ej. intenta OCR sin tesseract instalado -> modo básico
        pass
    # 2) Modo básico página por página
    return "\n\n".join(_pagina_a_md(p) for p in doc)


# ----------------------------------------------------------------- DOCX/DOC

def docx_a_md(ruta: Path) -> str:
    import docx
    from docx.document import Document as Doc
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(ruta)
    partes = []

    def iter_bloques(doc: Doc):
        from docx.oxml.ns import qn

        for hijo in doc.element.body.iterchildren():
            if hijo.tag == qn("w:p"):
                yield Paragraph(hijo, doc)
            elif hijo.tag == qn("w:tbl"):
                yield Table(hijo, doc)

    for bloque in iter_bloques(d):
        if isinstance(bloque, Paragraph):
            texto = bloque.text.strip()
            if texto:
                estilo = (bloque.style.name or "").lower()
                partes.append(f"## {texto}" if "heading" in estilo else texto)
        else:  # Table
            data = [[c.text for c in fila.cells] for fila in bloque.rows]
            md = tabla_a_md(data)
            if md:
                partes.append(md)
    return "\n\n".join(partes)


def doc_a_md(ruta: Path) -> str:
    """Convierte .doc antiguo a .docx con LibreOffice y luego a Markdown."""
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "docx",
             "--outdir", tmp, str(ruta)],
            check=True, capture_output=True, timeout=180,
        )
        convertido = next(Path(tmp).glob("*.docx"))
        return docx_a_md(convertido)


# --------------------------------------------------------------------- HTML

def html_a_md(ruta: Path) -> str:
    from bs4 import BeautifulSoup, UnicodeDammit
    from markdownify import ATX, MarkdownConverter

    # Read raw bytes and handle mixed encodings with detwingle
    raw_bytes = ruta.read_bytes()
    # Fix mixed encodings (common in gov.ec pages: UTF-8 + Windows-1252)
    fixed_bytes = UnicodeDammit.detwingle(raw_bytes)
    if fixed_bytes != raw_bytes:
        raw_bytes = fixed_bytes

    # Detect encoding
    dammit = UnicodeDammit(raw_bytes, is_html=True)
    encoding = dammit.original_encoding or "utf-8"

    # Parse with detected encoding
    soup = BeautifulSoup(raw_bytes, "lxml", from_encoding=encoding)

    # Extract links and images before decomposing for provenance
    links = []
    for a in soup.find_all("a", href=True):
        links.append({"url": a["href"], "text": a.get_text(strip=True)})

    images = []
    for img in soup.find_all("img", src=True):
        images.append({"src": img["src"], "alt": img.get("alt", "")})

    for tag in soup(["script", "style", "noscript", "svg", "form",
                     "nav", "footer", "iframe", "select", "button"]):
        tag.decompose()

    principal = soup.find("main") or soup.find("article") or soup.body or soup
    # Use MarkdownConverter for better table and image handling
    converter = MarkdownConverter(
        heading_style=ATX,
        table_infer_header=True,
        keep_inline_images_in=["p", "td"],
        bullets="*-+",
    )
    md = converter.convert(str(principal))

    lineas, vistas = [], set()
    for lin in md.split("\n"):
        lin = re.sub(r"[ \t]+", " ", lin).strip()
        if not lin:
            if lineas and lineas[-1] != "":
                lineas.append("")
            continue
        # Descartar líneas de navegación repetidas (menús de la página guardada)
        clave = lin.lower()
        if len(lin) < 45 and clave in vistas:
            continue
        vistas.add(clave)
        lineas.append(lin)
    return "\n".join(lineas).strip()


def txt_a_md(ruta: Path) -> str:
    return ruta.read_text(encoding="utf-8", errors="ignore").strip()


# ------------------------------------------------------------------ orquesta

def archivos_de_carpeta(carpeta: Path) -> list[Path]:
    """Lista los archivos de contenido de una carpeta de beca, en orden útil:
    primero la reseña/ficha, luego la página HTML y al final el resto."""
    archivos = [
        p for p in carpeta.rglob("*")
        if p.is_file()
        and p.suffix.lower() in EXTENSIONES
        and not any(PATRON_ASSETS.search(parte) for parte in p.parts)
    ]

    def prioridad(p: Path):
        nombre = p.name.lower()
        if "reseña" in nombre or "resena" in nombre:
            return (0, nombre)
        if p.suffix.lower() in (".html", ".htm"):
            return (1, nombre)
        return (2, nombre)

    return sorted(archivos, key=prioridad)


def extraer_carpeta(carpeta: Path) -> str:
    """Extrae todo el contenido de la carpeta a un solo Markdown consolidado."""
    secciones = []
    for ruta in archivos_de_carpeta(carpeta):
        ext = ruta.suffix.lower()
        try:
            if ext == ".pdf":
                md = pdf_a_md(ruta)
            elif ext == ".docx":
                md = docx_a_md(ruta)
            elif ext == ".doc":
                md = doc_a_md(ruta)
            elif ext in (".html", ".htm"):
                md = html_a_md(ruta)
            else:
                md = txt_a_md(ruta)
        except Exception as e:  # un archivo malo no detiene el proceso
            md = f"[No se pudo extraer {ruta.name}: {e}]"
        if md and md.strip():
            secciones.append(f"## Archivo: {ruta.name}\n\n{md.strip()}")
    return "\n\n".join(secciones)


def extract_folder(folder: Path) -> str:
    return extraer_carpeta(folder)
